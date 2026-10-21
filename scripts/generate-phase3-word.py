"""Generate current delivery Word files, preserving previous component guides."""
from copy import deepcopy
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / 'backend/docs'
FRONTEND = ROOT / 'frontend/climate-monitoring-web'
HISTORY = 'Anexo histórico — documentación anterior a Fase 3'


def markdown(doc, path):
    code = False
    table = None
    for line in path.read_text(encoding='utf-8-sig').splitlines():
        if line.startswith('```'):
            code = not code
            continue
        if not line.strip():
            continue
        text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', line).replace('**', '').replace('`', '')
        if code:
            p = doc.add_paragraph(text)
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(8)
        elif line.startswith('#'):
            doc.add_heading(text.lstrip('# '), min(len(line)-len(line.lstrip('#'))+1, 4))
        elif line.startswith('|'):
            if re.match(r'^\|[\s:|\-]+$', line):
                continue
            cells = [cell.strip() for cell in text.strip('|').split('|')]
            if table is None:
                table = doc.add_table(rows=0, cols=len(cells))
                table.style = 'Table Grid'
            for cell, value in zip(table.add_row().cells, cells):
                cell.text = value
            continue
        elif line.startswith('- '):
            doc.add_paragraph(text[2:], style='List Bullet')
        else:
            doc.add_paragraph(text)
        table = None


def previous_body(path):
    if not path.exists():
        return []
    original = Document(path)
    body = list(original.element.body)
    markers = [i for i, el in enumerate(body) if HISTORY in ''.join(el.itertext())]
    if markers:
        body = body[markers[0]+1:]
        # Explanatory paragraph belongs to the generated wrapper.
        if body and 'Se conserva como referencia' in ''.join(body[0].itertext()):
            body = body[1:]
    return [deepcopy(el) for el in body if not el.tag.endswith('}sectPr')]


def build(path, title, sources, preserve=False):
    history = previous_body(path) if preserve else []
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    doc.styles['Normal'].font.name = 'Aptos'
    doc.styles['Normal'].font.size = Pt(10)
    for level in range(1, 5):
        doc.styles[f'Heading {level}'].font.color.rgb = RGBColor(31, 78, 121)
    doc.add_heading(title, 0)
    doc.add_paragraph('Climate Monitoring System\nEntrega de Fase 3 — 21 de octubre de 2026')
    doc.add_paragraph('Estado: implementación funcional validada localmente. VPS y conservación de históricos excluidos del alcance acordado. Credenciales y archivos privados de entorno no forman parte de esta entrega.')
    doc.add_heading('Contenido de la entrega', 1)
    for source in sources:
        doc.add_paragraph(source.stem.replace('_', ' '), style='List Bullet')
    for source in sources:
        doc.add_page_break()
        markdown(doc, source)
    if history:
        doc.add_page_break()
        doc.add_heading(HISTORY, 1)
        doc.add_paragraph('Se conserva como referencia de construcción. Las versiones y decisiones vigentes se describen en la entrega de Fase 3 que precede a este anexo.')
        for el in history:
            doc.element.body.insert(len(doc.element.body)-1, el)
    section.header.paragraphs[0].text = 'Climate Monitoring System — Fase 3'
    doc.core_properties.title = title
    doc.core_properties.subject = 'Entrega local de Fase 3, arquitectura, uso y validación'
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    check = Document(path)
    text = '\n'.join(p.text for p in check.paragraphs)
    assert 'Fase 3' in text and 'simula' in text.lower()
    print(f'Validated {path.relative_to(ROOT)} ({path.stat().st_size} bytes)')


if __name__ == '__main__':
    sources = [BACKEND/name for name in ['FASE_3_INFORME_FINAL.md', 'FASE_3_ARQUITECTURA.md', 'FASE_3_REPORTERIA.md', 'FASE_3_DATOS_DEMO.md', 'FASE_3_PRUEBAS.md', 'FASE_3_FUERA_DE_ALCANCE.md']]
    frontend = [FRONTEND/'FASE_3_ENTREGA.md']
    build(BACKEND/'Documentacion_Esencial_Backend.docx', 'Documentación esencial del Backend — Fase 3', sources, True)
    build(FRONTEND/'docs/angular-components/Documentacion_Componentes_Angular.docx', 'Documentación de Componentes Angular — Fase 3', frontend, True)
    build(ROOT/'Documentacion_Climate_Monitoring_System.docx', 'Documentación técnica consolidada — Fase 3', sources+frontend)
